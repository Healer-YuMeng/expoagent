"""Document upload, parsing, chunking, and storage pipeline."""

from __future__ import annotations

import io
import logging
import uuid
import base64
from dataclasses import dataclass
from pathlib import Path
from typing import Optional

import pdfplumber
import docx
import pptx
import openpyxl
from openai import OpenAI
from pptx.enum.shapes import MSO_SHAPE_TYPE

from app.core.config import settings
from app.services.rag_chunking import build_parent_child_chunks
from app.services.rag_chunking import ensure_parent_chunk_metadata
from app.services.storage import StorageClient

logger = logging.getLogger(__name__)

SUPPORTED_EXTS = {".pdf", ".docx", ".pptx", ".xlsx", ".txt", ".md", ".jpg", ".jpeg", ".png"}
PPTX_IMAGE_MIME_TYPES = {
    "png": "image/png",
    "jpg": "image/jpeg",
    "jpeg": "image/jpeg",
    "webp": "image/webp",
}
OCR_EMPTY_MARKERS = {"", "(empty OCR)", "【Qwen OCR未配置】", "【Qwen OCR调用失败】"}


@dataclass
class DocMeta:
    doc_id: str
    name: str
    object_name: str
    size: int
    status: str
    chunks: int = 0
    school_id: str | None = None
    admin_id: str | None = None


class DocIngestor:
    """Handle file upload -> store -> parse -> chunk."""

    def __init__(self, storage: StorageClient, base_dir: Path) -> None:
        self.storage = storage
        self.base_dir = base_dir
        self.base_dir.mkdir(parents=True, exist_ok=True)

    def ingest(self, *, filename: str, content: bytes, content_type: Optional[str] = None, school_id: str | None = None, admin_id: str | None = None) -> tuple[DocMeta, list[dict]]:
        ext = Path(filename).suffix.lower()
        if ext not in SUPPORTED_EXTS:
            raise ValueError(f"Unsupported file type: {ext}")

        doc_id = str(uuid.uuid4())
        object_name = f"docs/{doc_id}{ext}"
        meta = DocMeta(doc_id=doc_id, name=filename, object_name=object_name, size=len(content), status="pending", school_id=school_id, admin_id=admin_id)

        # store file
        storage_uri = self.storage.put(object_name, io.BytesIO(content), length=len(content), content_type=content_type)
        logger.info("Stored doc %s at %s", doc_id, storage_uri)

        # parse content
        if ext == ".xlsx":
            rows = self._parse_xlsx_rows(content)
            chunks = self._chunk_xlsx_rows(rows, meta)
        else:
            text = self._parse(ext, content)
            chunks = self._chunk(text, meta)

        meta.status = "ready"
        meta.chunks = len(chunks)
        return meta, chunks

    def _ocr_image(self, content: bytes, mime_type: str) -> str:
        api_key = settings.QWEN_API_KEY
        api_base = settings.QWEN_API_BASE
        model = settings.RAG_IMAGE_OCR_MODEL or "qwen3-vl-plus"
        if not api_key or not api_base:
            return "【Qwen OCR未配置】"
        try:
            encoded = base64.b64encode(content).decode()
            client = OpenAI(api_key=api_key, base_url=api_base.rstrip("/"))
            completion = client.chat.completions.create(
                model=model,
                temperature=0,
                messages=[
                    {
                        "role": "system",
                        "content": "你是一个OCR文本提取助手。请准确提取图片中所有可见文字，尽量保留原有换行、段落和表格阅读顺序；不要补充解释；如果图片中没有可识别文字，只输出 (empty OCR)。",
                    },
                    {
                        "role": "user",
                        "content": [
                            {
                                "type": "image_url",
                                "image_url": {
                                    "url": f"data:{mime_type};base64,{encoded}",
                                },
                            },
                            {
                                "type": "text",
                                "text": "请提取这张图片中的全部文字内容。",
                            },
                        ],
                    },
                ],
            )
            content_text = completion.choices[0].message.content or ""
            if isinstance(content_text, str):
                normalized = content_text.strip()
                return normalized or "(empty OCR)"
            normalized = "".join(
                part.text for part in content_text if getattr(part, "type", "") == "text" and getattr(part, "text", "")
            ).strip()
            return normalized or "(empty OCR)"
        except Exception as exc:  # pragma: no cover
            logger.warning("Qwen OCR 调用失败，返回占位: %s", exc)
            return "【Qwen OCR调用失败】"

    def _extract_pptx_picture_text(self, shape: pptx.shapes.base.BaseShape) -> str | None:
        try:
            image = shape.image
        except Exception:  # pragma: no cover
            return None

        image_ext = (getattr(image, "ext", "") or "").lower().lstrip(".")
        mime_type = PPTX_IMAGE_MIME_TYPES.get(image_ext)
        if not mime_type:
            logger.info("Skip unsupported PPTX image type for OCR: %s", image_ext or "unknown")
            return None

        ocr_text = self._ocr_image(image.blob, mime_type).strip()
        if ocr_text in OCR_EMPTY_MARKERS:
            return None
        return ocr_text

    def _extract_pptx_shape_texts(self, shape: pptx.shapes.base.BaseShape) -> list[str]:
        texts: list[str] = []

        if getattr(shape, "shape_type", None) == MSO_SHAPE_TYPE.GROUP:
            for child in shape.shapes:
                texts.extend(self._extract_pptx_shape_texts(child))
            return texts

        raw_text = getattr(shape, "text", None)
        if isinstance(raw_text, str):
            normalized = raw_text.strip()
            if normalized:
                texts.append(normalized)

        if getattr(shape, "shape_type", None) == MSO_SHAPE_TYPE.PICTURE:
            picture_text = self._extract_pptx_picture_text(shape)
            if picture_text:
                texts.append(picture_text)

        return texts

    def _parse(self, ext: str, content: bytes) -> str:
        if ext in {".txt", ".md"}:
            return content.decode("utf-8", errors="ignore")
        if ext == ".pdf":
            with pdfplumber.open(io.BytesIO(content)) as pdf:
                return "\n".join(page.extract_text() or "" for page in pdf.pages)
        if ext == ".docx":
            doc = docx.Document(io.BytesIO(content))
            texts = [p.text for p in doc.paragraphs]
            for table in doc.tables:
                for row in table.rows:
                    texts.append(" ".join(cell.text for cell in row.cells))
            return "\n".join(texts)
        if ext == ".pptx":
            pres = pptx.Presentation(io.BytesIO(content))
            slides = []
            for slide in pres.slides:
                texts = []
                for shape in slide.shapes:
                    texts.extend(self._extract_pptx_shape_texts(shape))
                slide_text = "\n".join(texts).strip()
                if slide_text:
                    slides.append(slide_text)
            return "\n".join(slides)
        if ext in {".jpg", ".jpeg", ".png"}:
            mime_type = "image/png" if ext == ".png" else "image/jpeg"
            return self._ocr_image(content, mime_type)
        return ""

    def _parse_xlsx_rows(self, content: bytes) -> list[dict]:
        wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
        items: list[dict] = []
        for sheet in wb:
            rows = list(sheet.iter_rows(values_only=True))
            if not rows:
                continue
            headers = [str(h).strip() if h is not None else "" for h in rows[0]]
            for row_index, row in enumerate(rows[1:], start=1):
                pairs: list[str] = []
                for header, value in zip(headers, row):
                    if value is None:
                        continue
                    header_text = header.strip()
                    value_text = str(value).strip()
                    if not value_text:
                        continue
                    pairs.append(f"{header_text}:{value_text}" if header_text else value_text)
                row_text = " ".join(pairs).strip()
                if not row_text:
                    continue
                items.append(
                    {
                        "sheet_name": sheet.title,
                        "row_index": row_index,
                        "content": row_text,
                    }
                )
        return items

    def _chunk(self, text: str, meta: DocMeta) -> list[dict]:
        if not text.strip():
            return []
        base_metadata = {
            "doc_id": meta.doc_id,
            "doc_name": meta.name,
            "object_name": meta.object_name,
            "school_id": meta.school_id,
            "admin_id": meta.admin_id,
        }
        docs = build_parent_child_chunks(text, metadata=base_metadata, parent_seed=meta.doc_id)
        chunks: list[dict] = []
        for idx, d in enumerate(docs):
            metadata = dict(d.get("metadata") or {})
            metadata["chunk_index"] = idx
            chunks.append(
                {
                    "doc_id": meta.doc_id,
                    "doc_name": meta.name,
                    "object_name": meta.object_name,
                    "section_index": 0,
                    "chunk_index": idx,
                    "content": d.get("content", ""),
                    "metadata": metadata,
                }
            )
        return chunks

    def _chunk_xlsx_rows(self, rows: list[dict], meta: DocMeta) -> list[dict]:
        if not rows:
            return []

        chunks: list[dict] = []
        for idx, row in enumerate(rows):
            row_content = str(row.get("content") or "").strip()
            if not row_content:
                continue

            metadata = ensure_parent_chunk_metadata(
                row_content,
                {
                    "doc_id": meta.doc_id,
                    "doc_name": meta.name,
                    "object_name": meta.object_name,
                    "school_id": meta.school_id,
                    "admin_id": meta.admin_id,
                    "sheet_name": row.get("sheet_name"),
                    "row_index": row.get("row_index"),
                    "chunk_index": idx,
                },
                parent_seed=f"{meta.doc_id}-xlsx-row-{idx}",
            )
            chunks.append(
                {
                    "doc_id": meta.doc_id,
                    "doc_name": meta.name,
                    "object_name": meta.object_name,
                    "section_index": 0,
                    "chunk_index": idx,
                    "content": row_content,
                    "metadata": metadata,
                }
            )
        return chunks
